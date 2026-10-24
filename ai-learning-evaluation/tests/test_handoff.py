"""Persistent two-account review, restart, rollback and immutable exports."""
from contextlib import contextmanager
import copy
import sqlite3
import tkinter as tk
import pytest
from docx import Document
from src.ui.desktop import DesktopApp
from src.storage.reports import ReportRepository
from src.storage.snapshot import unpack
from src.ai.local_provider import LocalAnalysisProvider
from src.reporting.report_generator import generate_report
from test_account_session import session
from test_desktop import load_sample, decide_all_claims

def docx_text(path):
    """Paragraphs plus table rows (label: value) of the designed report."""
    doc = Document(path)
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            lines.append(': '.join(cells) if len(cells) == 2 and cells[0] and ':' not in cells[0] else ' '.join(cells))
    return '\n'.join(lines)



@contextmanager
def another_account(session, name):
    root = tk.Toplevel(session.root.master)
    app = DesktopApp(root, username=name, on_logout=lambda: None)
    app.repository = ReportRepository(session.repository.path)
    try:
        yield app
    finally:
        if not app.disposed:
            app.dispose_workspace()
        root.destroy()


def ready_report(app, audience='facilitator'):
    load_sample(app, count=12)
    app._apply_generated_report(generate_report(app.review_context(), audience, LocalAnalysisProvider()))
    decide_all_claims(app)
    app.library_title = 'Handoff demo'


def test_submit_switch_approve_restart_and_repeat_export(session, monkeypatch, tmp_path):
    ready_report(session)
    assert session.submit_for_approval()
    rid = session.library_id
    assert session.workflow.status == 'AWAITING HUMAN REVIEW'
    assert not session.library_unsaved()
    assert session.repository.list_pending('alex') == []
    assert [r['id'] for r in session.repository.list_pending('sam')] == [rid]
    assert not session.can_edit_workspace()
    assert session.editor.cget('state') == 'disabled'
    assert not session.approve_persistent()
    session.root.update()
    assert not session.library_unsaved()
    assert session.sign_out(True)
    with another_account(session, 'sam') as reviewer:
        reviewer.select_library_view('pending')
        assert rid in reviewer.library_cards
        assert reviewer.open_library_report(rid)
        assert reviewer.author_entry.get() == 'alex'
        assert reviewer.approver_entry.get() == 'sam'
        assert str(reviewer.approver_entry.cget('state')) == 'readonly'
        assert reviewer.can_review_workspace()
        assert reviewer.editor.cget('state') == 'disabled'
        assert reviewer.propose_button.instate(['disabled'])
        assert reviewer.submit_button.instate(['disabled'])
        old = reviewer.workspace_snapshot()
        reviewer.set_editor_text('Attempt to change the submitted report')
        reviewer.generate()
        reviewer.reject_claim()
        assert reviewer.workspace_snapshot() == old
        assert not reviewer.save_library_draft()
        reviewer.confirmed.set(True)
        assert reviewer.export()  # Approve report button: persist first, no file dialog.
        assert reviewer.workflow.approver == 'sam'
        assert reviewer.workflow.status == 'HUMAN REVIEWED'
        assert reviewer.repository.list_pending('sam') == []
        assert reviewer.export_button.cget('text') == 'Export approved'
        assert reviewer.sign_out()
    with another_account(session, 'alex') as owner:
        owner.select_library_view('approved')
        assert rid in owner.library_cards
        assert owner.open_library_report(rid)
        assert not owner.can_edit_workspace()
        assert owner.workflow.reviewer == 'alex'
        assert owner.workflow.approver == 'sam'
        monkeypatch.setattr('tkinter.filedialog.asksaveasfilename', lambda **kw: '')
        assert not owner.export()
        assert owner.workflow.status == 'HUMAN REVIEWED'
        md = tmp_path / 'approved.md'
        monkeypatch.setattr('tkinter.filedialog.asksaveasfilename', lambda **kw: str(md))
        assert owner.export()
        text = md.read_text()
        assert 'Prepared and submitted by: alex' in text
        assert 'Approved by: sam' in text
        assert 'different signed-in accounts' in text
        docx = tmp_path / 'approved.docx'
        monkeypatch.setattr('tkinter.filedialog.asksaveasfilename', lambda **kw: str(docx))
        assert owner.export()
        assert 'Approved by: sam' in docx_text(docx)
        assert owner.repository.load(rid, 'alex')['version'] == 3
    assert session.repository.list_approved('outsider') == []
    with pytest.raises(ValueError):
        session.repository.load(rid, 'outsider')


def test_return_notes_survive_restart_and_resubmission(session, monkeypatch):
    ready_report(session)
    assert session.submit_for_approval()
    rid = session.library_id
    with another_account(session, 'sam') as reviewer:
        assert reviewer.open_library_report(rid)
        monkeypatch.setattr('tkinter.simpledialog.askstring', lambda *a, **kw: '')
        assert not reviewer.return_for_changes()
        monkeypatch.setattr('tkinter.simpledialog.askstring', lambda *a, **kw: 'Clarify the delivery recommendation.')
        assert reviewer.return_for_changes()
        assert reviewer.current_page == 3
        assert reviewer.library_view == 'pending'
        assert reviewer.context is None
        assert reviewer.repository.list_pending('sam') == []
    # The original author still has the old submitted version open.
    assert session.open_library_report(rid)
    assert session.workflow.status == 'CHANGES REQUESTED'
    assert 'Clarify the delivery recommendation.' in session.return_notes.get()
    assert session.can_edit_workspace()
    assert session.save_library_draft()
    with another_account(session, 'alex') as restarted:
        assert restarted.open_library_report(rid)
        assert 'Clarify the delivery recommendation.' in restarted.return_notes.get()
        restarted.set_editor_text(restarted.editor.get('1.0', 'end-1c') + '\n\nWe recommend reviewing delivery next term.')
        restarted.refresh_claims()
        assert restarted.submit_for_approval()  # new wording is checked automatically
        assert restarted.workflow.times_returned() == 1
        assert len(restarted.repository.list_pending('sam')) == 1
    with another_account(session, 'sam') as reviewer:
        assert reviewer.open_library_report(rid)
        reviewer.confirmed.set(True)
        assert reviewer.approve_persistent()
        assert 'Times returned for changes: 1' in reviewer.approved_export


def test_pending_version_is_locked_and_review_conflicts_preserve_state(session, monkeypatch):
    ready_report(session)
    assert session.save_library_draft()
    before = session.workspace_snapshot()
    old_version = session.library_version
    assert session.submit_for_approval()
    rid, version = session.library_id, session.library_version
    for actor in ('alex', 'sam'):
        with pytest.raises(ValueError):
            session.repository.save(actor=actor, title='Overwrite', course='Course', status='DRAFT',
                snapshot=before, report_id=rid, version=version)
    for action in ('approve', 'return'):
        with pytest.raises(ValueError):
            session.repository.transition(actor='alex', report_id=rid, version=version,
                action=action, confirmed=True, notes='Return')
        with pytest.raises(ValueError):
            session.repository.transition(actor='sam', report_id=rid, version=old_version,
                action=action, confirmed=True, notes='Return')
    with another_account(session, 'sam') as first, another_account(session, 'lee') as stale:
        assert first.open_library_report(rid)
        assert stale.open_library_report(rid)
        first.confirmed.set(True)
        assert first.approve_persistent()
        stale.confirmed.set(True)
        errors = []
        monkeypatch.setattr('tkinter.messagebox.showerror', lambda *a, **kw: errors.append(a))
        assert not stale.approve_persistent()
        assert errors
        assert stale.workflow.status == 'AWAITING HUMAN REVIEW'
        assert unpack(session.repository.load(rid, 'alex')['snapshot'])['workflow'].approver == 'sam'


def test_confirmation_role_audit_failure_and_export_failure(session, monkeypatch, tmp_path):
    ready_report(session, audience='client')
    assert session.submit_for_approval()
    rid = session.library_id
    errors = []
    monkeypatch.setattr('tkinter.messagebox.showerror', lambda *a, **kw: errors.append(a))
    with another_account(session, 'sam') as reviewer:
        assert reviewer.open_library_report(rid)
        reviewer.approver_role.set('Facilitator')
        assert not reviewer.approve_persistent()
        reviewer.approver_role.set('Learning Futures lead')
        append = reviewer.audit.append
        monkeypatch.setattr(reviewer.audit, 'append', lambda *a, **kw: (_ for _ in ()).throw(OSError('disk full')))
        assert not reviewer.approve_persistent()
        assert reviewer.workflow.status == 'AWAITING HUMAN REVIEW'
        assert reviewer.repository.load(rid, 'sam')['status'] == 'AWAITING HUMAN REVIEW'
        monkeypatch.setattr(reviewer.audit, 'append', append)
        assert reviewer.approve_persistent()
        approved_version = reviewer.library_version
        monkeypatch.setattr('tkinter.filedialog.asksaveasfilename', lambda **kw: str(tmp_path / 'missing' / 'file.md'))
        assert not reviewer.export()
        assert reviewer.workflow.status == 'HUMAN REVIEWED'
        assert reviewer.repository.load(rid, 'sam')['version'] == approved_version
        assert errors


def test_repository_rejects_unreviewed_submission_and_other_owner(session):
    load_sample(session, count=12, review=False)
    session.library_title = 'Early draft'
    assert session.save_library_draft()
    for actor in ('alex', 'sam'):
        with pytest.raises(ValueError):
            session.repository.transition(actor=actor, report_id=session.library_id,
                version=session.library_version, action='submit')
    assert session.repository.load(session.library_id, 'alex')['status'] == 'DRAFT'
    assert session.repository.list_pending('sam') == []


def test_old_database_migrates_without_losing_drafts(tmp_path):
    path = tmp_path / 'old.sqlite3'
    with sqlite3.connect(path) as db:
        db.execute('CREATE TABLE reports (id TEXT PRIMARY KEY, owner TEXT NOT NULL, title TEXT NOT NULL, course TEXT NOT NULL, status TEXT NOT NULL, updated TEXT NOT NULL, version INTEGER NOT NULL, snapshot TEXT NOT NULL)')
        db.execute("INSERT INTO reports VALUES ('old','alex','Old draft','Course','DRAFT','2026-01-01',1,'{}')")
    repository = ReportRepository(path)
    assert repository.list_owned('alex')[0]['title'] == 'Old draft'
    assert repository.load('old', 'alex')['snapshot'] == {}
    assert repository.list_pending('sam') == []
    repository.save(actor='alex', title='Another', course='Course', status='DRAFT', snapshot={})
    assert len(repository.list_owned('alex')) == 2
