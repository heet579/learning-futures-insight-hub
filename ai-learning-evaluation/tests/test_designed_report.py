"""The two-page designed report: text tidying, emphasis, Word and PDF output."""
from io import BytesIO
from pathlib import Path

import pytest

from src.config import DATA_DIR

pytest.importorskip('matplotlib')
pytest.importorskip('reportlab')

from src import workspace  # noqa: E402
from src.ai.local_provider import LocalAnalysisProvider  # noqa: E402
from src.reporting import charts, designed_report as d  # noqa: E402

SAMPLE = Path(DATA_DIR) / 'synthetic_qualtrics_evaluation.csv'


@pytest.fixture(scope='module')
def state():
    s = workspace.new_state(workspace.load_files([SAMPLE]), [SAMPLE.name])
    for theme in s['context'].themes:
        workspace.decide_theme(s, theme.name, 'confirmed')
    workspace.generate(s, 'facilitator', provider=LocalAnalysisProvider())
    return s


def test_humanise_and_emphasis():
    assert d.humanise('**DRAFT — REQUIRES HUMAN REVIEW** 1 response(s) and 2781 related comments') == \
        '1 response and 2,781 related comments'
    parts = d.emphasis_runs('Averaged 4.08/5 with 85% recommending')
    assert ('4.08/5', True) in parts and ('85%', True) in parts and ('Averaged ', False) in parts


def test_report_data_reads_the_reviewed_draft(state):
    data = d.report_data(state)
    assert data['responses'] > 0 and data['working'] and data['recommendations']
    assert not data['approved'] and 'not yet approved' in d.status_line(data)
    assert len(d.figures(data)) == 4


def test_word_and_pdf_are_built_with_the_logo(state):
    from docx import Document
    data = d.report_data(state)
    word = Document(BytesIO(d.build_docx(data)))
    assert word.sections[0].header.tables[0].rows[0].cells[0].paragraphs[0].runs  # logo picture
    text = '\n'.join(p.text for p in word.paragraphs)
    assert 'At a glance' in text and 'About this report' in text and '**' not in text
    pypdf = pytest.importorskip('pypdf')
    assert len(pypdf.PdfReader(BytesIO(d.build_pdf(data))).pages) == 2


def test_described_charts_and_limits(state):
    assert workspace.describe_chart(state, 'average relevance by course', replace=0).startswith('Changed to:')
    assert state['charts'][0]['kind'] == 'group_bar' and state['charts'][0]['group'] == 'CourseName'
    with pytest.raises(ValueError):
        workspace.describe_chart(state, 'themes')      # already four charts
    with pytest.raises(ValueError):
        workspace.set_charts(state, [{'kind': 'ratings_bar'}] * 5)
    for i in range(len(state['charts'])):
        assert workspace.chart_png(state, i)[:4] == b'\x89PNG'
    assert charts.render({'kind': 'trend', 'measure': 'OverallSatisfaction', 'title': 't'},
                         state['frame'], state['context'])[:4] == b'\x89PNG'
