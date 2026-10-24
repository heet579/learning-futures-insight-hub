"""The two-page evaluation report (Word and PDF) with the Adelaide University logo and charts.

Page 1: who and what, key figures, the summary, two charts, what is working / what to improve.
Page 2: two more charts, what learners said, recommended next steps, notes, and the review record.

The text comes from the reviewed report; only its presentation changes here (plurals, thousands
separators and bold numbers). Every number and claim stays exactly as the reviewer approved it.
"""
from __future__ import annotations

import io
import re
from datetime import datetime
from pathlib import Path

from src.reporting import charts as chart_lib

LOGO = Path(__file__).resolve().parents[2] / 'assets' / 'adelaide-university-logo.png'
NAVY, PLUM, LILAC_BG, GREEN_BG, AMBER_BG = '151054', '5B2EA6', 'F1EDFA', 'E8F3EC', 'FCF1DE'
GREY, INK = '6B6382', '1E1631'
DRAFT_MARK = 'DRAFT — REQUIRES HUMAN REVIEW'


# ---- reading the reviewed text --------------------------------------------------------------------
def _sections(text):
    from src.review.claims import section_spans
    return {name: text[start:end].strip() for name, start, end in section_spans(text)}


def _bullets(body):
    items = []
    for line in body.splitlines():
        if re.match(r'^- ', line):
            items.append(line[2:].strip())
    return items


def humanise(text):
    """Presentation-only tidying: plurals, thousands separators, no markdown or status banners."""
    text = text.replace('**', '').replace(DRAFT_MARK, '').strip()
    text = re.sub(r'\b1 (response|comment|learner)\(s\)', r'1 \1', text)
    text = re.sub(r'\b(\w+)\(s\)', r'\1s', text)
    text = re.sub(r'\b(\d{4,})(?= (?:responses?|comments?|learners?|answers?|matching|related)\b)',
                  lambda m: f'{int(m.group(1)):,}', text)
    return re.sub(r'\s{2,}', ' ', text).strip()


_NUMBER = re.compile(r'(\d[\d,]*(?:\.\d+)?\s?(?:%|/5)?)')


def emphasis_runs(text):
    """Split text into (part, bold) pairs: numbers, percentages and x/5 scores are bold."""
    parts, last = [], 0
    for match in _NUMBER.finditer(text):
        if match.start() > last:
            parts.append((text[last:match.start()], False))
        parts.append((match.group(0), True))
        last = match.end()
    if last < len(text):
        parts.append((text[last:], False))
    return parts


def report_data(state, text=None):
    """Everything the two-page layout needs, from a saved workspace state."""
    from src.workspace import reviewed_context
    text = text if text is not None else (state.get('approved_export') or state['content'])
    approved = state['workflow'].status == 'HUMAN REVIEWED' and '## Review record' in text
    record = ''
    if '## Review record' in text:
        text, record = text.split('## Review record', 1)
    sections = _sections(text)
    context = reviewed_context(state)
    m = context.metrics
    audience = state['report'].audience if state.get('report') else 'facilitator'
    quotes = re.findall(r'Representative feedback:\s*[“"](.+?)[”"]', sections.get('Evidence / Supporting Metrics', ''))
    workflow = state['workflow']
    approved_event = next((e for e in reversed(workflow.events) if e.action == 'Approved'), None)
    return {
        'audience': audience,
        'kicker': 'Facilitator report' if audience == 'facilitator' else 'Client report',
        'title': context.course_name if context.course_name != 'All courses (aggregate)' else 'All courses',
        'source': state.get('loaded_file_label') or state.get('source', ''),
        'responses': m['response_count'], 'completeness': m['response_completeness'],
        'satisfaction': m['ratings']['OverallSatisfaction']['mean'], 'recommend': m.get('recommendation_percent'),
        'strongest': m['ratings'].get(m.get('strongest_area') or '', {}).get('label'),
        'lowest': m['ratings'].get(m.get('lowest_area') or '', {}).get('label'),
        'summary': humanise(sections.get('Executive Summary', '')),
        'working': [humanise(b) for b in _bullets(sections.get('What Worked Well', ''))],
        'improve': [humanise(b) for b in _bullets(sections.get('Areas for Improvement', ''))],
        'recommendations': [humanise(b) for b in _bullets(sections.get('Recommendations', ''))],
        'participation': humanise(sections.get('Participation Overview', '')),
        'method': humanise(sections.get('Analysis method', '')),
        'quotes': quotes[:2] if audience == 'facilitator' else [],
        'record': [humanise(line[2:]) for line in record.splitlines() if line.startswith('- ')],
        'approved': approved,
        'approver': workflow.approver if approved else '',
        'approved_at': approved_event.timestamp if approved and approved_event else '',
        'author': workflow.reviewer,
        'context': context, 'frame': state['frame'],
        'charts': chart_lib.chart_list(state.get('charts') or chart_lib.DEFAULT_CHARTS),
    }


def _date(iso=''):
    try:
        moment = datetime.fromisoformat(iso).astimezone() if iso else datetime.now().astimezone()
    except ValueError:
        moment = datetime.now().astimezone()
    return f'{moment.day} {moment:%B %Y}'


def status_line(data):
    if data['approved']:
        return f"Approved by {data['approver']} on {_date(data['approved_at'])}"
    return 'Draft for internal review: not yet approved'


def figures(data):
    """The key-figure strip: (value, label)."""
    items = [(f"{data['responses']:,}", 'responses analysed')]
    if data['satisfaction'] is not None:
        items.append((f"{data['satisfaction']:.2f}", 'average satisfaction out of 5'))
    if data['recommend'] is not None:
        items.append((f"{data['recommend']:.0f}%", 'would recommend the course'))
    items.append((f"{data['completeness']:.0f}%", 'answers complete'))
    return items[:4]


CHART_W, CHART_H = 3.55, 3.05


def chart_images(data, width=CHART_W, height=CHART_H):
    return [(spec['title'], chart_lib.render(spec, data['frame'], data['context'], width, height))
            for spec in data['charts']]


# ---- Word --------------------------------------------------------------------------------------
def build_docx(data, logo=LOGO):
    from docx import Document
    from docx.enum.section import WD_SECTION_START  # noqa: F401
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(1.7)
    section.top_margin, section.bottom_margin = Cm(2.3), Cm(1.6)
    section.header_distance, section.footer_distance = Cm(0.8), Cm(0.7)

    base = doc.styles['Normal']
    base.font.name, base.font.size = 'Calibri', Pt(9.5)
    base.font.color.rgb = RGBColor.from_string(INK)
    base.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    base.paragraph_format.space_after = Pt(3)
    base.paragraph_format.line_spacing = 1.08

    def shade(cell, colour):
        tc = cell._tc.get_or_add_tcPr()
        fill = OxmlElement('w:shd')
        fill.set(qn('w:val'), 'clear'); fill.set(qn('w:color'), 'auto'); fill.set(qn('w:fill'), colour)
        tc.append(fill)

    def borders(cell, **sides):
        tc = cell._tc.get_or_add_tcPr()
        node = OxmlElement('w:tcBorders')
        for side in ('top', 'left', 'bottom', 'right'):
            spec = sides.get(side)
            el = OxmlElement(f'w:{side}')
            if spec:
                el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), str(spec[0])); el.set(qn('w:color'), spec[1])
            else:
                el.set(qn('w:val'), 'nil')
            node.append(el)
        tc.append(node)

    def margins(cell, top=90, side=140):
        tc = cell._tc.get_or_add_tcPr()
        node = OxmlElement('w:tcMar')
        for edge, value in (('top', top), ('bottom', top), ('left', side), ('right', side)):
            el = OxmlElement(f'w:{edge}'); el.set(qn('w:w'), str(value)); el.set(qn('w:type'), 'dxa'); node.append(el)
        tc.append(node)

    def no_table_borders(table):
        tbl = table._tbl.tblPr
        node = OxmlElement('w:tblBorders')
        for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            el = OxmlElement(f'w:{side}'); el.set(qn('w:val'), 'nil'); node.append(el)
        tbl.append(node)

    def run(paragraph, text, size=None, bold=False, colour=None, italic=False, font=None):
        r = paragraph.add_run(text)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
        if colour:
            r.font.color.rgb = RGBColor.from_string(colour)
        if font:
            r.font.name = font
            r._element.rPr.rFonts.set(qn('w:eastAsia'), font)
        return r

    def rich(paragraph, text, size=None, colour=None, lead_bold=False):
        if lead_bold and ':' in text[:60]:
            head, tail = text.split(':', 1)
            run(paragraph, head + ':', size, True, colour)
            text = tail
        for part, bold in emphasis_runs(text):
            run(paragraph, part, size, bold, colour)

    def heading(container, text, size=12, colour=NAVY, before=6, after=3):
        p = container.add_paragraph()
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
        p.paragraph_format.keep_with_next = True
        run(p, text, size, True, colour, font='Georgia')
        return p

    def bullet_list(cell, items, colour=INK, empty='Nothing recorded.'):
        for item in items or [empty]:
            p = cell.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Cm(0.45), Cm(-0.35)
            rich(p, item, 9, colour, lead_bold=True)

    def field(paragraph, code):
        for kind, text in (('begin', None), (None, code), ('separate', None), (None, '1'), ('end', None)):
            r = paragraph.add_run()
            if kind:
                el = OxmlElement('w:fldChar'); el.set(qn('w:fldCharType'), kind); r._r.append(el)
            elif text == code:
                el = OxmlElement('w:instrText'); el.set(qn('xml:space'), 'preserve'); el.text = f' {code} '; r._r.append(el)
            else:
                r.text = text
            r.font.size, r.font.color.rgb = Pt(7.5), RGBColor.from_string(GREY)

    # header: logo left, report name right, a thin navy rule
    header = section.header
    table = header.add_table(1, 2, Cm(17.6))
    no_table_borders(table)
    left, right = table.rows[0].cells
    left.width, right.width = Cm(8.8), Cm(8.8)
    if Path(logo).is_file():
        left.paragraphs[0].add_run().add_picture(str(logo), height=Cm(1.05))
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(p, 'Learning Futures\n', 8, True, NAVY)
    run(p, 'Course evaluation report', 8, False, GREY)
    header.paragraphs[0].paragraph_format.space_after = Pt(0)
    rule = header.add_paragraph()
    pPr = rule._p.get_or_add_pPr()
    bdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    for key, value in (('val', 'single'), ('sz', '8'), ('space', '1'), ('color', NAVY)):
        bottom.set(qn(f'w:{key}'), value)
    bdr.append(bottom); pPr.append(bdr)
    rule.paragraph_format.space_after = Pt(0)

    # footer: status and date left, page x of y right
    foot = section.footer.add_table(1, 2, Cm(17.6))
    no_table_borders(foot)
    foot_left, foot_right = foot.rows[0].cells
    foot_left.width, foot_right.width = Cm(12.6), Cm(5)
    run(foot_left.paragraphs[0], f"{status_line(data)}  ·  Prepared {_date()}", 7.5, False, GREY)
    footer = foot_right.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(footer, 'Page ', 7.5, False, GREY)
    field(footer, 'PAGE')
    run(footer, ' of ', 7.5, False, GREY)
    field(footer, 'NUMPAGES')

    # title block
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    run(p, data['kicker'], 9, True, PLUM)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    run(p, data['title'], 20, True, NAVY, font='Georgia')
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run(p, f"{data['responses']:,} survey responses  ·  {data['source']}  ·  {status_line(data)}", 8.5, False, GREY)

    # key figures
    items = figures(data)
    strip = doc.add_table(1, len(items))
    strip.alignment = WD_TABLE_ALIGNMENT.CENTER
    no_table_borders(strip)
    for cell, (value, label) in zip(strip.rows[0].cells, items):
        shade(cell, LILAC_BG); margins(cell, 110, 160)
        borders(cell, left=(18, NAVY))
        p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        run(p, value, 17, True, NAVY, font='Georgia')
        p = cell.add_paragraph(); p.paragraph_format.space_after = Pt(0)
        run(p, label, 8, False, GREY)

    # at a glance
    heading(doc, 'At a glance', before=10)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    rich(p, data['summary'] or 'No summary was written.', 9.5)
    highlights = []
    if data['strongest']:
        highlights.append(f"Strongest area: {data['strongest']}")
    if data['lowest'] and data['lowest'] != data['strongest']:
        highlights.append(f"Most room to improve: {data['lowest']}")
    if highlights:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        for i, item in enumerate(highlights):
            head, tail = item.split(':', 1)
            run(p, ('    ' if i else '') + head + ':', 9, True, PLUM)
            run(p, tail, 9, True, INK)

    images = chart_images(data)

    def chart_row(pairs, start):
        if not pairs:
            return
        row = doc.add_table(1, 2)
        no_table_borders(row)
        for offset, cell in enumerate(row.rows[0].cells):
            cell.width = Cm(8.8)
            if offset >= len(pairs):
                continue
            title, png = pairs[offset]
            cell.paragraphs[0].add_run().add_picture(io.BytesIO(png), width=Cm(8.5))
            caption = cell.add_paragraph()
            caption.paragraph_format.space_after = Pt(2)
            run(caption, f'Figure {start + offset}. ', 8, True, NAVY)
            run(caption, title, 8, False, GREY, italic=True)

    chart_row(images[:2], 1)

    # what is working / what to improve
    boxes = doc.add_table(1, 2)
    no_table_borders(boxes)
    for cell, (title, items, fill, accent, empty) in zip(boxes.rows[0].cells, (
            ("What's working", data['working'], GREEN_BG, '2E7D5B', 'No recurring strength stood out.'),
            ('What to improve', data['improve'], AMBER_BG, 'C27C0E', 'No recurring concern stood out.'))):
        shade(cell, fill); margins(cell, 100, 160); borders(cell, top=(18, accent))
        p = cell.paragraphs[0]
        run(p, title, 10.5, True, NAVY, font='Georgia')
        bullet_list(cell, items[:5], empty=empty)

    # page 2
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    chart_row(images[2:4], 3)

    if data['quotes']:
        heading(doc, 'In their words', before=6)
        for quote in data['quotes']:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.4)
            p.paragraph_format.space_after = Pt(3)
            pPr = p._p.get_or_add_pPr(); bdr = OxmlElement('w:pBdr'); left_b = OxmlElement('w:left')
            for key, value in (('val', 'single'), ('sz', '18'), ('space', '8'), ('color', PLUM)):
                left_b.set(qn(f'w:{key}'), value)
            bdr.append(left_b); pPr.append(bdr)
            run(p, f'“{quote}”', 9.5, False, INK, italic=True)

    # recommended next steps, in a navy box
    box = doc.add_table(1, 1)
    no_table_borders(box)
    cell = box.rows[0].cells[0]
    shade(cell, NAVY); margins(cell, 140, 200)
    p = cell.paragraphs[0]
    run(p, 'Recommended next steps', 12, True, 'FFFFFF', font='Georgia')
    for number, item in enumerate(data['recommendations'] or ['No recommendation was kept.'], 1):
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Cm(0.55), Cm(-0.55)
        run(p, f'{number}.  ', 9.5, True, 'C9B8EE')
        rich(p, item, 9.5, 'FFFFFF', lead_bold=True)
    p = cell.add_paragraph()
    run(p, 'Recommendations are proposals for the team to discuss, not proven fixes.', 8, False, 'C9B8EE', italic=True)

    # notes and review record
    heading(doc, 'About this report', size=10.5, before=10)
    for text in (data['participation'], data['method'],
                 'Ratings are learner perceptions, not measured learning outcomes. Theme counts are comments, '
                 'not people, and a comment can touch several themes.'):
        if text:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            rich(p, text, 8.5, GREY)
    heading(doc, 'Review record' if data['approved'] else 'Review status', size=10.5, before=8)
    if data['record']:
        rec = doc.add_table(0, 2)
        no_table_borders(rec)
        for line in data['record'][:9]:
            if ':' not in line:
                continue
            label, value = line.split(':', 1)
            value = value.strip()
            if re.fullmatch(r'[0-9a-f]{64}.*', value):
                value = value[:16] + '…' + value[64:]
            cells = rec.add_row().cells
            cells[0].width, cells[1].width = Cm(6), Cm(11.6)
            run(cells[0].paragraphs[0], label.strip(), 8, True, NAVY)
            run(cells[1].paragraphs[0], value, 8, False, INK)
    else:
        p = doc.add_paragraph()
        rich(p, 'This draft has not been approved yet. A second person must review every theme, statement and '
                'quote before it is shared.', 8.5, GREY)

    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


# ---- PDF -----------------------------------------------------------------------------------------
def build_pdf(data, logo=LOGO):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (BaseDocTemplate, Frame, Image, KeepInFrame, NextPageTemplate, PageBreak,
                                    PageTemplate, Paragraph, Spacer, Table, TableStyle)

    navy, plum, grey, ink = (colors.HexColor('#' + c) for c in (NAVY, PLUM, GREY, INK))
    lilac, green, amber = (colors.HexColor('#' + c) for c in (LILAC_BG, GREEN_BG, AMBER_BG))
    body = ParagraphStyle('body', fontName='Helvetica', fontSize=9.2, leading=12.2, textColor=ink, spaceAfter=3)
    small = ParagraphStyle('small', parent=body, fontSize=8.2, leading=10.6, textColor=grey)
    h = ParagraphStyle('h', fontName='Times-Bold', fontSize=12.5, leading=15, textColor=navy, spaceBefore=7, spaceAfter=3)
    h_small = ParagraphStyle('hs', parent=h, fontSize=10.5, leading=13)
    bullet = ParagraphStyle('b', parent=body, fontSize=8.8, leading=11.4, leftIndent=9, bulletIndent=0, spaceAfter=2)

    def markup(text, lead_bold=False):
        from xml.sax.saxutils import escape
        head = ''
        if lead_bold and ':' in text[:60]:
            first, text = text.split(':', 1)
            head = f'<b>{escape(first)}:</b>'
        return head + ''.join(f'<b>{escape(p)}</b>' if b else escape(p) for p, b in emphasis_runs(text))

    width, height = A4
    margin = 1.7 * cm

    def decorate(canvas, doc):
        canvas.saveState()
        if Path(logo).is_file():
            canvas.drawImage(str(logo), margin, height - 1.95 * cm, height=1.05 * cm, width=1.05 * cm * 800 / 190,
                             mask='auto')
        canvas.setFont('Helvetica-Bold', 8); canvas.setFillColor(navy)
        canvas.drawRightString(width - margin, height - 1.3 * cm, 'Learning Futures')
        canvas.setFont('Helvetica', 8); canvas.setFillColor(grey)
        canvas.drawRightString(width - margin, height - 1.65 * cm, 'Course evaluation report')
        canvas.setStrokeColor(navy); canvas.setLineWidth(0.8)
        canvas.line(margin, height - 2.1 * cm, width - margin, height - 2.1 * cm)
        canvas.setFont('Helvetica', 7.5); canvas.setFillColor(grey)
        canvas.drawString(margin, 0.9 * cm, f"{status_line(data)}  ·  Prepared {_date()}")
        canvas.drawRightString(width - margin, 0.9 * cm, f'Page {doc.page} of 2')
        canvas.restoreState()

    frame = Frame(margin, 1.5 * cm, width - 2 * margin, height - 3.9 * cm, id='main', leftPadding=0, rightPadding=0,
                  topPadding=0, bottomPadding=0)
    stream = io.BytesIO()
    doc = BaseDocTemplate(stream, pagesize=A4, title=f"{data['kicker']}: {data['title']}", author='Learning Futures',
                          leftMargin=margin, rightMargin=margin)
    doc.addPageTemplates([PageTemplate('page', [frame], onPage=decorate)])
    inner = width - 2 * margin
    images = chart_images(data)

    def chart_row(pairs, start):
        if not pairs:
            return []
        cells = []
        for offset, (title, png) in enumerate(pairs):
            img = Image(io.BytesIO(png), width=inner / 2 - 0.25 * cm, height=(inner / 2 - 0.25 * cm) * CHART_H / CHART_W)
            caption = Paragraph(f'<font color="#{NAVY}"><b>Figure {start + offset}.</b></font> <i>{title}</i>', small)
            cells.append([img, caption])
        while len(cells) < 2:
            cells.append('')
        table = Table([cells], colWidths=[inner / 2] * 2)
        table.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'TOP'), ('LEFTPADDING', (0, 0), (-1, -1), 0)]))
        return [table, Spacer(1, 4)]

    def box(title, items, fill, accent, empty):
        flow = [Paragraph(title, ParagraphStyle('bt', parent=h_small, spaceBefore=0))]
        flow += [Paragraph(markup(i, True), bullet, bulletText='•') for i in (items[:5] or [empty])]
        return flow

    page1 = [
        Paragraph(f'<font color="#{PLUM}"><b>{data["kicker"]}</b></font>', body),
        Paragraph(data['title'], ParagraphStyle('t', fontName='Times-Bold', fontSize=21, leading=24, textColor=navy)),
        Paragraph(f"{data['responses']:,} survey responses  ·  {data['source']}  ·  {status_line(data)}", small),
        Spacer(1, 6),
    ]
    stats = [[Paragraph(f'<font name="Times-Bold" size="17" color="#{NAVY}">{v}</font><br/>'
                        f'<font size="7.8" color="#{GREY}">{l}</font>', ParagraphStyle('s', leading=14)) for v, l in figures(data)]]
    strip = Table(stats, colWidths=[inner / len(stats[0])] * len(stats[0]))
    strip.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, -1), lilac), ('LINEBEFORE', (0, 0), (-1, -1), 2, navy),
                               ('TOPPADDING', (0, 0), (-1, -1), 7), ('BOTTOMPADDING', (0, 0), (-1, -1), 7),
                               ('LEFTPADDING', (0, 0), (-1, -1), 9)]))
    page1 += [strip, Paragraph('At a glance', h), Paragraph(markup(data['summary'] or 'No summary was written.'), body)]
    highlights = []
    if data['strongest']:
        highlights.append(f"<font color='#{PLUM}'><b>Strongest area:</b></font> <b>{data['strongest']}</b>")
    if data['lowest'] and data['lowest'] != data['strongest']:
        highlights.append(f"<font color='#{PLUM}'><b>Most room to improve:</b></font> <b>{data['lowest']}</b>")
    if highlights:
        page1.append(Paragraph('&nbsp;&nbsp;&nbsp;&nbsp;'.join(highlights), body))
    page1 += chart_row(images[:2], 1)
    boxes = Table([[box("What's working", data['working'], green, None, 'No recurring strength stood out.'),
                    box('What to improve', data['improve'], amber, None, 'No recurring concern stood out.')]],
                  colWidths=[inner / 2 - 3, inner / 2 - 3])
    boxes.setStyle(TableStyle([('BACKGROUND', (0, 0), (0, 0), green), ('BACKGROUND', (1, 0), (1, 0), amber),
                               ('LINEABOVE', (0, 0), (0, 0), 2, colors.HexColor('#2E7D5B')),
                               ('LINEABOVE', (1, 0), (1, 0), 2, colors.HexColor('#C27C0E')),
                               ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('TOPPADDING', (0, 0), (-1, -1), 7),
                               ('LEFTPADDING', (0, 0), (-1, -1), 9), ('RIGHTPADDING', (0, 0), (-1, -1), 9)]))
    page1.append(boxes)

    page2 = chart_row(images[2:4], 3)
    if data['quotes']:
        page2.append(Paragraph('In their words', h))
        quote_style = ParagraphStyle('q', parent=body, fontName='Helvetica-Oblique', leftIndent=10, borderPadding=(2, 0, 2, 6))
        for quote in data['quotes']:
            from xml.sax.saxutils import escape
            q = Table([[Paragraph(f'“{escape(quote)}”', quote_style)]], colWidths=[inner])
            q.setStyle(TableStyle([('LINEBEFORE', (0, 0), (0, 0), 2, plum), ('LEFTPADDING', (0, 0), (-1, -1), 8)]))
            page2 += [q, Spacer(1, 3)]
    white = ParagraphStyle('w', parent=body, textColor=colors.white, fontSize=9.3, leading=12.4, leftIndent=14,
                           bulletIndent=0, spaceAfter=3)
    recs = [Paragraph('Recommended next steps', ParagraphStyle('rh', parent=h, textColor=colors.white, spaceBefore=0))]
    recs += [Paragraph(markup(item, True), white, bulletText=f'{n}.')
             for n, item in enumerate(data['recommendations'] or ['No recommendation was kept.'], 1)]
    recs.append(Paragraph('<i>Recommendations are proposals for the team to discuss, not proven fixes.</i>',
                          ParagraphStyle('rn', parent=small, textColor=colors.HexColor('#C9B8EE'))))
    rec_box = Table([[recs]], colWidths=[inner])
    rec_box.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, -1), navy), ('TOPPADDING', (0, 0), (-1, -1), 10),
                                 ('BOTTOMPADDING', (0, 0), (-1, -1), 10), ('LEFTPADDING', (0, 0), (-1, -1), 13),
                                 ('RIGHTPADDING', (0, 0), (-1, -1), 13)]))
    page2 += [Spacer(1, 6), rec_box, Paragraph('About this report', h_small)]
    for text in (data['participation'], data['method'], 'Ratings are learner perceptions, not measured learning '
                 'outcomes. Theme counts are comments, not people, and a comment can touch several themes.'):
        if text:
            page2.append(Paragraph(markup(text), small))
    page2.append(Paragraph('Review record' if data['approved'] else 'Review status', h_small))
    if data['record']:
        rows = []
        for line in data['record'][:9]:
            if ':' not in line:
                continue
            label, value = line.split(':', 1)
            value = value.strip()
            if re.fullmatch(r'[0-9a-f]{64}.*', value):
                value = value[:16] + '…' + value[64:]
            rows.append([Paragraph(f'<b><font color="#{NAVY}">{label.strip()}</font></b>', small), Paragraph(value, small)])
        if rows:
            record = Table(rows, colWidths=[5.6 * cm, inner - 5.6 * cm])
            record.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'TOP'), ('LEFTPADDING', (0, 0), (-1, -1), 0),
                                        ('TOPPADDING', (0, 0), (-1, -1), 1), ('BOTTOMPADDING', (0, 0), (-1, -1), 1)]))
            page2.append(record)
    else:
        page2.append(Paragraph('This draft has not been approved yet. A second person must review every theme, '
                               'statement and quote before it is shared.', small))

    available = (inner, height - 3.9 * cm)
    # Each page is shrunk to fit if needed, so the PDF is always exactly two pages.
    doc.build([KeepInFrame(*available, page1, mode='shrink'), NextPageTemplate('page'), PageBreak(),
               KeepInFrame(*available, page2, mode='shrink')])
    return stream.getvalue()


MEDIA = {'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
         'pdf': 'application/pdf', 'md': 'text/markdown; charset=utf-8'}


def export_bytes(state, text, fmt):
    """The designed two-page report (docx/pdf) or plain Markdown (md) for this text."""
    if fmt == 'md':
        from src.reporting.exporter import markdown_bytes
        return markdown_bytes(text)
    if fmt not in ('docx', 'pdf'):
        raise ValueError('Choose Word (.docx), PDF (.pdf) or Markdown (.md).')
    data = report_data(state, text)
    return build_docx(data) if fmt == 'docx' else build_pdf(data)
