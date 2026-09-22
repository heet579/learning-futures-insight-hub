from io import BytesIO
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

def markdown_bytes(content: str) -> bytes:
    return content.encode("utf-8")

def docx_bytes(content: str, logo_path: str | None = None) -> bytes:
    """Compact single-column layout: tight margins/spacing so a typical report fits on one page,
    with a header reserved for a university logo (a real image if `logo_path` is given)."""
    doc = Document()

    section = doc.sections[0]
    section.left_margin = section.right_margin = Inches(0.6)
    section.top_margin = section.bottom_margin = Inches(0.5)

    normal = doc.styles["Normal"]
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(4)

    heading1 = doc.styles["Heading 1"]
    heading1.font.size = Pt(14)
    heading1.paragraph_format.space_before = Pt(0)
    heading1.paragraph_format.space_after = Pt(6)

    heading2 = doc.styles["Heading 2"]
    heading2.font.size = Pt(10.5)
    heading2.paragraph_format.space_before = Pt(8)
    heading2.paragraph_format.space_after = Pt(2)

    bullet = doc.styles["List Bullet"]
    bullet.font.size = Pt(9.5)
    bullet.paragraph_format.space_after = Pt(2)

    header_paragraph = section.header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if logo_path and Path(logo_path).is_file():
        header_paragraph.add_run().add_picture(logo_path, height=Inches(0.4))
    else:
        run = header_paragraph.add_run("[University logo]")
        run.italic = True
        run.font.size = Pt(8)

    for line in content.splitlines():
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line.strip():
            doc.add_paragraph(line.replace("**", ""))
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()
